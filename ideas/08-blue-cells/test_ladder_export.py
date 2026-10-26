#!/usr/bin/env python3
"""Tests for ladder_export.py (unittest). From the bank root:
    python3 -m unittest discover -s ideas/08-blue-cells -p 'test_*.py'
External pins (README "Tests"): rev_a.json through the workbook's own formulas, DECISION_LOG items 6-8,
decision_grid.json's md5, the sidecar, skyhook_finance.py's row map, blob id 46bc02e. The ladder values
and census counts are regression pins read off the sheet XML.
"""
import contextlib
import functools
import hashlib
import importlib.util
import io
import json
import math
import os
import re
import shutil
import struct
import subprocess
import sys
import unittest
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.dont_write_bytecode = True
sys.path.insert(0, HERE)
import ladder_export as LE  # noqa: E402

SNAPS = os.path.join(HERE, "snapshots")
XLSX = os.path.join(SNAPS, "skyhook_financial_model.xlsx")
ROWS_MAP = os.path.join(SNAPS, "ladder_rows.json")
SIDECAR = XLSX + ".provenance.json"
LADDER_JSON = os.path.join(SNAPS, "ladder.json")
FIXTURE = os.path.join(HERE, "fixtures", "shuffled_assumptions.xlsx")
RECALC = os.path.join(HERE, "fixtures", "recalculated_by_libreoffice.xlsx")
PRE_FIX = os.path.join(HERE, "fixtures", "pre_fix_7bb5b5f.xlsx")
SKYHOOK = "/home/user/skyhook"


def _load(path):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def setUpModule():
    global WB, ROWS, RESULT, REV_A, GRID, PROV, LOG, XLSX_BYTES
    with open(XLSX, "rb") as fh:
        XLSX_BYTES = fh.read()
    ROWS = _load(ROWS_MAP)
    PROV = _load(SIDECAR)
    WB = LE.read_workbook(XLSX)
    RESULT = LE.analyze(WB, ROWS, PROV)
    REV_A = _load(os.path.join(SNAPS, "rev_a.json"))
    GRID = _load(os.path.join(SNAPS, "decision_grid.json"))
    with open(os.path.join(SNAPS, "DECISION_LOG.md"), encoding="utf-8") as fh:
        LOG = fh.read()


def tearDownModule():
    # plain `python3 -m unittest` would otherwise leave a __pycache__ here
    shutil.rmtree(os.path.join(HERE, "__pycache__"), ignore_errors=True)


SCN = ("Conservative", "Baseline", "Optimistic")

# Regression table: Assumptions!B3:D16 at Skyhook 34935cb.
EXPECTED_LADDER = {
    "strength_retention": (0.7, 0.85, 1.0),
    "ascent_losses_ms": (1800, 1500, 1200),
    "construction_launch_per_kg": (2500, 1500, 800),
    "power_chain_efficiency": (0.15, 0.2, 0.3),
    "electricity_per_kwh": (0.08, 0.06, 0.04),
    "ops_wear_per_kg": (12, 8, 5),
    "hardware_mass_factor": (1.8, 1.6, 1.4),
    "customer_price_per_kg": (500, 500, 500),
    "catches_per_day": (1, 2, 3),
    "payload_kg": (5000, 5000, 5000),
    "revenue_ramp_y1": (0.3, 0.4, 0.5),
    "fixed_ops_M": (80, 60, 45),
    "tip_speed_kms": (2.5, 2.5, 2.5),
    "program_fleet_capex_M": (6000, 5000, 4000),
}


def ladder_value(key, scn):
    return RESULT["ladder"]["rows"][key]["values"][scn]


def design(key):
    return RESULT["design_point"]["entries"][key]["value"]


def constant(key):
    return RESULT["constants"]["entries"][key]["value"]


def workbook_chain(scn, years=15, inputs=None, design=None, constant=None):
    """The workbook's formulas, transcribed from sheet3.xml (Design Calcs rows 3-4, 8-9, 11-13, 19-25)
    and sheet5.xml (Cash Flow rows 4-16), on the extracted inputs or on the given inputs/design/constant."""
    get = inputs or ladder_value
    design = design or globals()["design"]
    constant = constant or globals()["constant"]
    ret = get("strength_retention", scn)
    v_tip = get("tip_speed_kms", scn)
    payload = get("payload_kg", scn)
    launch = get("construction_launch_per_kg", scn)
    hw = get("hardware_mass_factor", scn)
    fleet = get("program_fleet_capex_M", scn)
    eff = get("power_chain_efficiency", scn)
    kwh = get("electricity_per_kwh", scn)
    ops = get("ops_wear_per_kg", scn)
    cad = get("catches_per_day", scn)
    price = get("customer_price_per_kg", scn)
    ramp = get("revenue_ramp_y1", scn)
    fixed = get("fixed_ops_M", scn)
    R = constant("earth_radius_km") + design("com_altitude_km")                       # Design Calcs B3
    v_orb = math.sqrt(constant("mu_earth_km3_s2") / R)                                # B4, EQ 2
    Vc = math.sqrt(2 * design("fiber_tensile_strength_sigma_gpa") * 1e9 * ret
                   / design("fiber_density_rho_kg_m3")) / 1000                        # B8, EQ 4 with derate
    x = v_tip / (Vc / math.sqrt(design("safety_factor_on_stress")))                   # B9
    single = math.sqrt(math.pi) * x * math.exp(x * x) * math.erf(x)                   # B11 Moravec/Forward
    tons = 2 * single * payload / 1000                                                 # B12
    catch = v_orb - v_tip - design("equatorial_rotation_help_km_s")                    # B13, EQ 3 + Earth help
    sys_capex = tons * 1000 * launch * hw / 1e6                                        # B19
    reboost = 2 * v_orb * v_tip                                                        # B20
    marginal = reboost / eff / 3.6 * kwh + ops                                         # B21 + B22
    upmass = cad * 365 * payload / 1000                                                # B23
    rev = upmass * 1000 * price / 1e6                                                  # B24
    mcost = upmass * 1000 * marginal / 1e6                                             # B25
    total = sys_capex + fleet                                                          # Cash Flow B4
    cum, payback, cumulative, net = 0.0, None, [], []
    for y in range(1, years + 1):
        spend = total * (0.15, 0.3, 0.35, 0.2)[y - 1] if y <= 4 else 0.0               # row 11
        if y >= 5:
            m = ramp if y == 5 else (min(1.0, ramp + 0.3) if y == 6 else 1.0)            # rows 12-13
            r, o = rev * m, mcost * m + fixed
        else:
            r, o = 0.0, 0.0
        net.append(r - spend - o)                                                      # row 14
        cum += net[-1]                                                                 # row 15
        cumulative.append(cum)
        if payback is None and cum > 0:                                                # row 16
            payback = y
    # NPV is not in the workbook; skyhook_finance.npv_M discounts year t's net flow by (1+rate)^t.
    npv10 = sum(cf / 1.10 ** t for t, cf in enumerate(net, start=1))
    return {"ratio": 2 * single, "tons": tons, "catch": catch, "total_capex": total, "payback": payback,
            "cumulative": cumulative, "npv10": npv10, "v_orb": v_orb, "R": R}


class TestLadderValues(unittest.TestCase):
    def test_every_ladder_cell(self):
        rows = RESULT["ladder"]["rows"]
        self.assertEqual(list(rows), list(EXPECTED_LADDER))
        for key, expected in EXPECTED_LADDER.items():
            for scn, want in zip(SCN, expected):
                self.assertEqual(rows[key]["values"][scn], want, f"{key} {scn}")
                self.assertEqual(rows[key]["cells"][scn], f"{ROWS['columns'][scn]}{ROWS['rows'][key]}")

    def test_rows_map_is_the_finance_module_s(self):
        self.assertEqual(ROWS["source"], "finance/skyhook_finance.py")
        self.assertEqual(ROWS["scenarios"], list(SCN))
        self.assertEqual(ROWS["columns"], {"Conservative": "B", "Baseline": "C", "Optimistic": "D"})
        self.assertEqual(sorted(ROWS["rows"].values()), list(range(3, 17)))
        self.assertEqual(RESULT["ladder"]["header_row"], 2)
        self.assertEqual(RESULT["ladder"]["sheet"], "Assumptions")
        self.assertEqual(set(LE.LABEL_GUARD), set(ROWS["rows"]), "every key has a drift-guard fragment")

    def test_raw_text_is_the_sheet_s(self):
        r = RESULT["ladder"]["rows"]["strength_retention"]["raw"]
        self.assertEqual(r, {"Conservative": "0.7", "Baseline": "0.85", "Optimistic": "1"})

    def test_notes_and_labels(self):
        rows = RESULT["ladder"]["rows"]
        self.assertIn("rev_a.json derate_fraction 30%/15%/0%", rows["strength_retention"]["note"])
        self.assertEqual(rows["construction_launch_per_kg"]["label"], "Launch price during construction ($/kg)")
        self.assertIn("fixed at Rev A 2.5 km/s reference", rows["tip_speed_kms"]["label"])

    def test_design_point_and_constants(self):
        d = RESULT["design_point"]["entries"]
        self.assertEqual([v["cell"] for v in d.values()], [f"B{r}" for r in range(19, 29)])
        self.assertEqual(d["com_altitude_km"]["value"], 610)
        self.assertEqual(d["half_length_l_km"]["value"], 320)
        self.assertEqual(d["safety_factor_on_stress"]["value"], 2.5)
        self.assertEqual(d["fiber_tensile_strength_sigma_gpa"]["value"], 5.8)
        self.assertEqual(d["fiber_density_rho_kg_m3"]["value"], 1560)
        self.assertEqual(d["equatorial_rotation_help_km_s"]["value"], 0.465)
        c = RESULT["constants"]["entries"]
        self.assertEqual(c, {
            "mu_earth_km3_s2": {"row": 32, "label": "mu Earth (km3/s2)", "cell": "B32", "value": 398600.0,
                                "raw": "398600", "note": None, "blue": True},
            "earth_radius_km": {"row": 33, "label": "Earth radius (km)", "cell": "B33", "value": 6371.0,
                                "raw": "6371", "note": None, "blue": True},
        })


class TestReconciliation(unittest.TestCase):
    """Pins to numbers published outside this folder's code."""

    def test_derate_ladder_matches_rev_a(self):
        mr = REV_A["mass_ratio"]["integrated_tether_to_payload_both_arms"]
        pct = {k: int(re.search(r"_(\d+)pct_derate$", k).group(1)) for k in mr if re.search(r"_(\d+)pct_derate$", k)}
        want = {"Conservative": REV_A["materials"]["derate_fraction"],
                "Baseline": pct["baseline_15pct_derate"] / 100,
                "Optimistic": pct["optimistic_0pct_derate"] / 100}
        self.assertEqual(want["Conservative"], 0.3)
        for scn in SCN:
            derate = 1 - ladder_value("strength_retention", scn)
            self.assertAlmostEqual(derate, want[scn], places=9, msg=scn)
        self.assertAlmostEqual(1 - ladder_value("strength_retention", "Conservative"),
                               pct["conservative_30pct_derate"] / 100, places=9)

    def test_tip_speed_is_the_rev_a_reference_in_every_column(self):
        for scn in SCN:
            self.assertEqual(ladder_value("tip_speed_kms", scn), REV_A["kinematics"]["v_tip_kms"])
        self.assertEqual(design("rev_a_reference_tip_speed_km_s"), REV_A["kinematics"]["v_tip_kms"])

    def test_design_point_matches_rev_a_materials_and_orbit(self):
        m = REV_A["materials"]
        self.assertEqual(design("fiber_tensile_strength_sigma_gpa"), m["tensile_strength_GPa"])
        self.assertEqual(design("fiber_density_rho_kg_m3"), m["density_kgm3"])
        self.assertEqual(design("safety_factor_on_stress"), m["safety_factor"])
        self.assertEqual(design("com_altitude_km"), REV_A["orbital"]["com_altitude_km"])
        self.assertEqual(design("half_length_l_km"), REV_A["geometry"]["tether_half_length_km"])
        self.assertEqual(ladder_value("payload_kg", "Conservative"),
                         REV_A["mass_ratio"]["integrated_tether_to_payload_both_arms"]["payload_mass_kg_for_tonnage"])

    def test_eq6_from_extracted_inputs_reproduces_rev_a_mass_ratios(self):
        mr = REV_A["mass_ratio"]["integrated_tether_to_payload_both_arms"]
        want = {"Conservative": ("conservative_30pct_derate", "conservative_30pct_fiber_mass_t"),
                "Baseline": ("baseline_15pct_derate", "baseline_15pct_fiber_mass_t"),
                "Optimistic": ("optimistic_0pct_derate", "optimistic_0pct_fiber_mass_t")}
        for scn, (k_ratio, k_tons) in want.items():
            got = workbook_chain(scn)
            self.assertAlmostEqual(got["ratio"], mr[k_ratio], places=2, msg=scn)
            self.assertAlmostEqual(got["tons"], mr[k_tons], places=1, msg=scn)

    def test_lower_tip_groundrel_matches_rev_a(self):
        got = workbook_chain("Conservative")["catch"]
        self.assertAlmostEqual(round(got, 2), REV_A["kinematics"]["v_lower_tip_groundrel_kms"], places=9)

    def test_capex_and_payback_reproduce_decision_log_item_6(self):
        m = re.search(r"Post-fix workbook results: Conservative ([\d.]+)× / (\d+) t / \$([\d,]+)M / payback ~Y(\d+); "
                      r"Baseline ([\d.]+)× / (\d+) t / \$([\d,]+)M / Y(\d+); "
                      r"Optimistic ([\d.]+)× / (\d+) t / \$([\d,]+)M / Y(\d+)", LOG)
        self.assertIsNotNone(m, "DECISION_LOG excerpt still states the post-fix results")
        g = m.groups()
        published = {
            "Conservative": (float(g[0]), int(g[1]), int(g[2].replace(",", "")), int(g[3])),
            "Baseline": (float(g[4]), int(g[5]), int(g[6].replace(",", "")), int(g[7])),
            "Optimistic": (float(g[8]), int(g[9]), int(g[10].replace(",", "")), int(g[11])),
        }
        for scn, (ratio, tons, capex, payback) in published.items():
            got = workbook_chain(scn)
            self.assertAlmostEqual(got["ratio"], ratio, places=2, msg=scn)
            self.assertEqual(round(got["tons"]), tons, scn)
            self.assertEqual(round(got["total_capex"]), capex, scn)
            if scn == "Conservative":
                self.assertIsNone(got["payback"], "the workbook's own 15-year window says 'beyond Y15'")
                self.assertEqual(workbook_chain(scn, years=40)["payback"], payback, "~Y17 by steady-state extrapolation")
            else:
                self.assertEqual(got["payback"], payback, scn)

    def test_conservative_npv_reproduces_decision_log_item_8(self):
        """DECISION_LOG.md:57 'NPV@10% -$3,883M over 15 years'. A 3% change in any Conservative cost row
        or the ramp changes the rounded figure, so this pin is tight."""
        m = re.search(r"conservative case \*\*fails the 15-year window\*\*: payback ~Y(\d+), NPV@10% −\$([\d,]+)M over 15 years", LOG)
        self.assertIsNotNone(m, "DECISION_LOG excerpt still states the conservative NPV (item 8, line 57)")
        published = -int(m.group(2).replace(",", ""))
        got = workbook_chain("Conservative")
        self.assertEqual(round(got["npv10"]), published)
        self.assertEqual(workbook_chain("Conservative", years=40)["payback"], int(m.group(1)))
        for key in ("electricity_per_kwh", "power_chain_efficiency", "ops_wear_per_kg", "fixed_ops_M", "revenue_ramp_y1"):
            nudged = workbook_chain("Conservative", inputs=lambda k, s, key=key: ladder_value(k, s) * (1.03 if k == key else 1))
            self.assertNotEqual(round(nudged["npv10"]), published, f"the NPV pin constrains {key}")

    def test_the_payback_years_are_loose_pins(self):
        """README 'How tight those pins are': on a 0.01 grid, Y9 and Y7 survive these single-input
        scalings (inclusive) and change one step beyond each end."""
        ranges = {  # (scenario, key): (lowest factor that keeps the year, highest)
            ("Baseline", "revenue_ramp_y1"): (0.13, 1.33),
            ("Baseline", "customer_price_per_kg"): (0.84, 1.08),
            ("Baseline", "catches_per_day"): (0.83, 1.08),
            ("Baseline", "power_chain_efficiency"): (0.04, None),
            ("Baseline", "fixed_ops_M"): (None, 5.14),
            ("Optimistic", "revenue_ramp_y1"): (0.32, 1.29),
            ("Optimistic", "customer_price_per_kg"): (0.71, 1.22),
            ("Optimistic", "catches_per_day"): (0.71, 1.22),
        }
        want = {"Baseline": 9, "Optimistic": 7}

        def payback(scn, key, f):
            return workbook_chain(scn, inputs=lambda k, s: ladder_value(k, s) * (f if k == key else 1))["payback"]

        for (scn, key), (lo, hi) in ranges.items():
            with self.subTest(scn=scn, key=key):
                for f, keeps in ((lo, True), (hi, True), (lo and round(lo - 0.01, 2), False),
                                 (hi and round(hi + 0.01, 2), False)):
                    if f:
                        self.assertEqual(payback(scn, key, f) == want[scn], keeps, f"x{f}")
        for scn in want:
            for key in ("electricity_per_kwh", "ops_wear_per_kg"):
                for f in (0.01, 10):
                    self.assertEqual(payback(scn, key, f), want[scn], f"{scn} {key} x{f}")

    def test_earth_radius_convention(self):
        """B33 = 6371 (physics/skyhook_physics.py:40), so R = 6981 km; rev_a.json's com_radius_km is 6988
        (R_E 6378). The 4.59 km/s lower-tip pin cannot tell the two apart."""
        self.assertEqual(constant("earth_radius_km"), 6371)
        R = workbook_chain("Conservative")["R"]
        self.assertEqual(R, 6981)
        self.assertNotEqual(R, REV_A["orbital"]["com_radius_km"])
        self.assertEqual(REV_A["orbital"]["com_radius_km"] - REV_A["orbital"]["com_altitude_km"], 6378)
        other = workbook_chain("Conservative", constant=lambda k: 6378.0 if k == "earth_radius_km" else constant(k))
        self.assertEqual(round(other["catch"], 2), round(workbook_chain("Conservative")["catch"], 2),
                         "both radii round to rev_a.json's 4.59, so that pin does not choose between them")
        self.assertAlmostEqual(other["v_orb"], REV_A["orbital"]["v_orb_kms"], places=2)

    def test_conservative_launch_price_is_the_one_decision_log_names(self):
        m = re.search(r"conservative column's own launch price is \$([\d,]+)/kg \(workbook Assumptions (B\d+)\)", LOG)
        self.assertIsNotNone(m)
        price, cell = int(m.group(1).replace(",", "")), m.group(2)
        row = RESULT["ladder"]["rows"]["construction_launch_per_kg"]
        self.assertEqual(row["cells"]["Conservative"], cell)
        self.assertEqual(row["values"]["Conservative"], price)
        self.assertGreater(price, row["values"]["Optimistic"], "the $800 column is the optimistic one")

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is not installed; the side-by-side read is skipped")
    def test_agrees_with_openpyxl_the_way_skyhook_finance_reads_it(self):
        """The 42 cells as load_ladder_from_workbook reads them (read_only, data_only=False) equal
        this reader's; with data_only=True openpyxl gives None for all 146 formula cells."""
        import openpyxl
        wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=False)
        ws = wb[ROWS["sheet"]]
        for key, row in ROWS["rows"].items():
            for scn, col in ROWS["columns"].items():
                self.assertEqual(float(ws[f"{col}{row}"].value), ladder_value(key, scn), f"{col}{row} {key}")
        for entry in list(RESULT["design_point"]["entries"].values()) + list(RESULT["constants"]["entries"].values()):
            self.assertEqual(float(ws[entry["cell"]].value), entry["value"], entry["cell"])
        wb.close()
        formulas = openpyxl.load_workbook(XLSX, data_only=False)
        values = openpyxl.load_workbook(XLSX, data_only=True)
        per_sheet = {}
        for ws_f, ws_v in zip(formulas.worksheets, values.worksheets):
            n_formula = n_cached = 0
            for row in ws_f.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        n_formula += 1
                        if ws_v[c.coordinate].value is not None:
                            n_cached += 1
            per_sheet[ws_f.title] = (n_formula, n_cached)
        self.assertEqual(per_sheet, {name: (c["formula_cells"], c["formula_cells_with_cached_value"])
                                     for name, c in RESULT["census"]["sheets"].items()})
        self.assertEqual(sum(f for f, _ in per_sheet.values()), 146)
        self.assertEqual(sum(v for _, v in per_sheet.values()), 0)

    def test_md5_closes_the_loop_with_decision_grid(self):
        self.assertEqual(RESULT["workbook"]["md5"], GRID["workbook_md5"])
        self.assertEqual(RESULT["workbook"]["sha256"], PROV["sourceSha256"])
        self.assertEqual(RESULT["workbook"]["bytes"], PROV["sourceBytes"])
        self.assertTrue(PROV["readMatchesCommitted"])
        self.assertEqual(RESULT["workbook"]["md5"], LE.md5_hex(XLSX_BYTES))


class TestCensus(unittest.TestCase):
    """Regression pins for what the file stores at 34935cb (81/63/2, 0 cached; the judges agree)."""

    def test_per_sheet_counts(self):
        cs = RESULT["census"]["sheets"]
        self.assertEqual(list(cs), ["README", "Assumptions", "Design Calcs", "Phase Budgets", "Cash Flow & Payback"])
        want = {  # formulas, cached, literal numbers, strings, blue, green
            "README": (0, 0, 0, 6, 0, 0),
            "Assumptions": (0, 0, 54, 52, 54, 0),
            "Design Calcs": (81, 0, 0, 58, 0, 0),
            "Phase Budgets": (2, 0, 8, 14, 8, 0),
            "Cash Flow & Payback": (63, 0, 19, 34, 1, 5),
        }
        for name, (f, cached, nums, strs, blue, green) in want.items():
            c = cs[name]
            self.assertEqual((c["formula_cells"], c["formula_cells_with_cached_value"], c["literal_numbers"],
                              c["strings"], c["blue_cells"], c["green_cells"]), (f, cached, nums, strs, blue, green), name)
            self.assertEqual(c["cached_refs"], [], name)
            self.assertEqual(c["blue_cells_with_formula"], [], name)
            self.assertEqual(c["green_cells_without_cross_sheet_formula"], [], name)

    def test_totals_and_finding(self):
        t = RESULT["census"]["totals"]
        self.assertEqual(t["formula_cells"], 146)
        self.assertEqual(t["formula_cells_with_cached_value"], 0)
        self.assertEqual(t["literal_numbers"], 81)
        self.assertEqual(t["blue_cells"], 63)
        # 54 on Assumptions + 8 on Phase Budgets are numbers; the 63rd blue cell is the scenario
        # selector Cash Flow & Payback!B2, a string ("Conservative").
        self.assertEqual(t["blue_literal_numbers"], 62)
        selector = next(c for c in WB["sheets"][4]["cells"] if c["ref"] == "B2")
        self.assertEqual((selector["kind"], selector["value"], selector["font_rgb"]), ("string", "Conservative", "0000FF"))
        self.assertEqual(t["blue_cells_with_formula"], 0)
        self.assertEqual(t["green_cells"], 5)
        self.assertEqual(t["yellow_fill_cells"], 4)
        self.assertEqual(t["cells"], 395)
        self.assertIn("146 formula cells and 0 cached results", RESULT["finding"])
        self.assertIn("fullCalcOnLoad=1", RESULT["finding"])
        self.assertNotIn("wrong", RESULT["finding"].lower())

    def test_every_formula_cell_has_an_empty_v(self):
        for s in WB["sheets"]:
            for c in s["cells"]:
                if c["kind"] == "formula":
                    self.assertIsNone(c["cached"], f"{s['name']}!{c['ref']}")
                    self.assertTrue(c["formula"], f"{s['name']}!{c['ref']} has a formula body")

    def test_equation_labels_predate_eq6(self):
        """Finding 7: README!A3, A5 and Design Calcs!A1, E20 number the equations as before EQ-6/EQ-7,
        now and before the fix; physics/skyhook_physics.py:13-16 is read when the checkout is
        present. A relabelled workbook fails this test and retires the finding."""
        stale = {
            ("README", "A3"): "Design Calcs (physics, EQ 1-5)",
            ("README", "A5"): "taper & Moravec/Forward mass ratio (EQ 4), ED reboost energy (EQ 5)",
            ("Design Calcs", "A1"): "all formulas, per scenario (EQ 1-5)",
            ("Design Calcs", "E20"): "EQ 5 momentum view: 2 v_orb v_tip",
            ("Design Calcs", "E8"): "EQ 4 with ladder derate",
        }
        for wb in (WB, LE.read_workbook(PRE_FIX)):
            grid = {(s["name"], c["ref"]): c for s in wb["sheets"] for c in s["cells"]}
            for where, text in stale.items():
                self.assertIn(text, grid[where]["value"], f"{wb['file']} {where}")
            self.assertTrue(grid[("Design Calcs", "A8")]["value"].startswith("Derated char. velocity V_c"))
            self.assertEqual(grid[("Design Calcs", "B11")]["formula"], "SQRT(PI())*B9*EXP(B9^2)*ERF(B9)")
        physics = os.path.join(SKYHOOK, "physics", "skyhook_physics.py")
        if not os.path.isfile(physics):
            self.skipTest(f"{physics} is not present; the canonical numbering is cited, not read")
        with open(physics, encoding="utf-8") as fh:
            lines = fh.read().split("\n")
        self.assertTrue(lines[12].strip().startswith("EQ-4  Characteristic velocity"), lines[12])
        self.assertTrue(lines[13].strip().startswith("EQ-5  Taper ratio"), lines[13])
        self.assertTrue(lines[14].strip().startswith("EQ-6  Integrated tether mass ratio"), lines[14])
        self.assertTrue(lines[15].strip().startswith("EQ-7  Upper-tip release energy"), lines[15])

    def test_yellow_cells_are_the_weakest_assumptions(self):
        cs = RESULT["census"]["sheets"]
        self.assertEqual(cs["Assumptions"]["yellow_fill_cells"], ["B8", "C8", "D8"])
        self.assertEqual(cs["Cash Flow & Payback"]["yellow_fill_cells"], ["B2"])

    def test_workbook_metadata(self):
        w = RESULT["workbook"]
        self.assertEqual(w["application"], "Microsoft Excel Compatible / Openpyxl 3.1.5")
        self.assertEqual(w["creator"], "openpyxl")
        self.assertEqual(w["calc_pr"]["fullCalcOnLoad"], "1")
        self.assertIsNone(w["shared_strings"])
        self.assertEqual(w["members"], 17)
        self.assertEqual(w["media"], [{"member": "xl/media/image1.png", "bytes": 53976,
                                       "md5": "11384969e42362346381174c57ddc7ae"}])
        self.assertTrue(w["legend"].startswith("Color code: BLUE = hardcoded inputs"))
        self.assertEqual([c["ok"] for c in RESULT["checks"]], [True] * 4)
        self.assertEqual([c["name"] for c in RESULT["checks"]], ["ladder-blue", "no-blue-formula", "green-are-links", "design-point"])
        self.assertEqual([g["name"] for g in RESULT["guards"]], ["ladder-rows", "ladder-literals"])

    def test_styles_read_the_legend_colours(self):
        grid = {(s["name"], c["ref"]): c for s in WB["sheets"] for c in s["cells"]}
        self.assertEqual(grid[("Assumptions", "B3")]["font_rgb"], "0000FF")
        self.assertEqual(grid[("Assumptions", "B8")]["fill_rgb"], "FFFF00")
        self.assertIsNone(grid[("Assumptions", "B3")]["fill_rgb"])
        self.assertEqual(grid[("Design Calcs", "B3")]["font_rgb"], "000000")
        self.assertEqual(grid[("Cash Flow & Payback", "B4")]["font_rgb"], "008000")
        self.assertEqual(grid[("Cash Flow & Payback", "B2")]["fill_rgb"], "FFFF00")
        self.assertEqual(grid[("Cash Flow & Payback", "F11")]["kind"], "number")  # a literal 0, not empty
        self.assertEqual(grid[("Cash Flow & Payback", "F11")]["value"], 0.0)
        self.assertEqual(grid[("Design Calcs", "E15")]["kind"], "empty")  # <c t="n"/> with no <v>


SHEET_NS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'


class TestParser(unittest.TestCase):
    def test_cell_kinds_and_entities(self):
        xml = (f'<worksheet {SHEET_NS}><sheetData><row r="1">'
               '<c r="A1" t="inlineStr"><is><t>Ops &amp; wear &gt; 5 &#x2014; &quot;x&quot;</t></is></c>'
               '<c r="B1" t="n"/>'
               '<c r="C1"><f>1+1</f><v/></c>'
               '<c r="D1"><f>IF(B1&gt;0,1,2)</f><v>2</v></c>'
               '<c r="E1" t="str"><f>"a"</f><v>a</v></c>'
               '<c r="F1" t="b"><v>1</v></c>'
               '<c r="G1" t="e"><v>#DIV/0!</v></c>'
               '<c r="H1"><v>0</v></c>'
               '<c r="I1" s="3"><v> 2.5 </v></c>'
               '<c r="J1" t="inlineStr"><is><r><t>two </t></r><r><t>runs</t></r></is></c>'
               '</row></sheetData></worksheet>').encode()
        cells = {c["ref"]: c for c in LE.parse_sheet(xml, None, None)}
        self.assertEqual(cells["A1"]["value"], 'Ops & wear > 5 — "x"')
        self.assertEqual(cells["B1"]["kind"], "empty")
        self.assertEqual((cells["C1"]["kind"], cells["C1"]["formula"], cells["C1"]["cached"]), ("formula", "1+1", None))
        self.assertEqual((cells["D1"]["formula"], cells["D1"]["cached"]), ("IF(B1>0,1,2)", "2"))
        self.assertEqual((cells["E1"]["kind"], cells["E1"]["cached"]), ("formula", "a"))
        self.assertEqual((cells["F1"]["kind"], cells["F1"]["value"]), ("bool", True))
        self.assertEqual((cells["G1"]["kind"], cells["G1"]["value"]), ("error", "#DIV/0!"))
        self.assertEqual((cells["H1"]["kind"], cells["H1"]["value"]), ("number", 0.0))
        self.assertEqual((cells["I1"]["value"], cells["I1"]["raw"], cells["I1"]["s"]), (2.5, "2.5", 3))
        self.assertEqual(cells["J1"]["value"], "two runs")

    def test_shared_strings_and_cached_values_in_a_synthetic_workbook(self):
        labels = {k: LE.LABEL_GUARD[k] for k in EXPECTED_LADDER}
        strings = ["Parameter", "Conservative", "Baseline", "Optimistic"] + [labels[k] for k in EXPECTED_LADDER]
        sst = (f'<sst {SHEET_NS} count="{len(strings)}" uniqueCount="{len(strings)}">'
               + "".join(f"<si><t>{s}</t></si>" for s in strings) + "</sst>")
        rows = ['<row r="2"><c r="A2" t="s"><v>0</v></c><c r="B2" t="s"><v>1</v></c>'
                '<c r="C2" t="s"><v>2</v></c><c r="D2" t="s"><v>3</v></c></row>']
        for i, (k, vals) in enumerate(EXPECTED_LADDER.items()):
            r = i + 3
            rows.append(f'<row r="{r}"><c r="A{r}" t="s"><v>{4 + i}</v></c>'
                        + "".join(f'<c r="{col}{r}"><v>{v}</v></c>' for col, v in zip("BCD", vals))
                        + (f'<c r="F{r}"><f>B{r}*2</f><v>{vals[0] * 2}</v></c><c r="G{r}"><f>B{r}*3</f><v/></c>' if i == 0 else "")
                        + "</row>")
        sheet = f'<worksheet {SHEET_NS}><sheetData>{"".join(rows)}</sheetData></worksheet>'
        wbxml = ('<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                 '<sheets><sheet name="Assumptions" sheetId="1" r:id="rId1"/></sheets><calcPr calcId="1"/></workbook>')
        rels = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="x" Target="worksheets/sheet1.xml"/></Relationships>')
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("xl/workbook.xml", wbxml, compress_type=zipfile.ZIP_STORED)
            z.writestr("xl/_rels/workbook.xml.rels", rels, compress_type=zipfile.ZIP_DEFLATED)
            z.writestr("xl/sharedStrings.xml", sst, compress_type=zipfile.ZIP_DEFLATED)
            z.writestr("xl/worksheets/sheet1.xml", sheet, compress_type=zipfile.ZIP_STORED)
        wb = LE.read_workbook(buf.getvalue(), name="synthetic.xlsx")
        self.assertEqual(wb["shared_strings"], len(strings))
        res = LE.analyze(wb, ROWS)
        for k, vals in EXPECTED_LADDER.items():
            self.assertEqual([res["ladder"]["rows"][k]["values"][s] for s in SCN], [float(v) for v in vals], k)
        self.assertFalse(res["ladder"]["rows"]["strength_retention"]["blue"], "no styles part: nothing is blue")
        c = res["census"]["sheets"]["Assumptions"]
        self.assertEqual((c["formula_cells"], c["formula_cells_with_cached_value"], c["cached_refs"]), (2, 1, ["F3"]))
        self.assertIn("2 formula cells, 1 of them with a cached result stored in the file; the other 1 store none",
                      res["finding"])
        checks = {c["name"]: c["ok"] for c in res["checks"]}
        self.assertFalse(checks["ladder-blue"], "ladder-blue reports the missing input styling")
        self.assertEqual(res["design_point"]["entries"], {})
        self.assertFalse(checks["design-point"])

    def test_any_relationships_prefix_and_single_quotes(self):
        """The relationship id is found by namespace, not the literal 'r:', and single quotes are XML
        too (bluecells.test.mjs runs the same XML)."""
        wbxml = ("<x:workbook xmlns:x='http://schemas.openxmlformats.org/spreadsheetml/2006/main' "
                 "xmlns:rel='http://schemas.openxmlformats.org/officeDocument/2006/relationships'>"
                 "<x:sheets><x:sheet name='Only' sheetId='1' rel:id='rId7'></x:sheet></x:sheets></x:workbook>")
        rels = ("<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>"
                "<Relationship Id='rId7' Type='x' Target='/xl/worksheets/one.xml'/></Relationships>")
        sheet = ("<x:worksheet xmlns:x='http://schemas.openxmlformats.org/spreadsheetml/2006/main'><x:sheetData>"
                 "<x:row r='1'><x:c r='A1' t='inlineStr'><x:is><x:t>hi</x:t></x:is></x:c><x:c r='B1'><x:v>7</x:v></x:c>"
                 "</x:row></x:sheetData></x:worksheet>")
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("xl/workbook.xml", wbxml)
            z.writestr("xl/_rels/workbook.xml.rels", rels)
            z.writestr("xl/worksheets/one.xml", sheet)
        wb = LE.read_workbook(buf.getvalue(), name="prefixed.xlsx")
        self.assertEqual([(s["name"], s["part"]) for s in wb["sheets"]], [("Only", "xl/worksheets/one.xml")])
        self.assertEqual([(c["ref"], c["kind"], c["value"]) for c in wb["sheets"][0]["cells"]],
                         [("A1", "string", "hi"), ("B1", "number", 7.0)])

    def test_missing_shared_strings_part_is_a_workbook_error(self):
        xml = f'<worksheet {SHEET_NS}><sheetData><row r="1"><c r="A1" t="s"><v>0</v></c></row></sheetData></worksheet>'.encode()
        with self.assertRaises(LE.WorkbookError):
            LE.parse_sheet(xml, None, None)

    def test_slug_and_refs(self):
        self.assertEqual(LE.slug("CoM altitude (km)"), "com_altitude_km")
        self.assertEqual(LE.slug("Rev A reference tip speed (km/s) — GTO/high-elliptical"), "rev_a_reference_tip_speed_km_s")
        self.assertEqual(LE.slug("Fiber strength retention (1 - derate); joints/weave/aging"), "fiber_strength_retention_1_derate")
        self.assertEqual(LE.split_ref("AB12"), ("AB", 12))
        self.assertEqual(LE.col_index("AB"), 28)
        with self.assertRaises(ValueError):
            LE.split_ref("12A")


def _swap_rows_in_snapshot(a, b):
    """The snapshot with Assumptions rows a and b exchanged, in memory, chart kept."""
    sys.path.insert(0, HERE)
    import build_data  # noqa: E402
    with zipfile.ZipFile(io.BytesIO(XLSX_BYTES)) as zin:
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for name in zin.namelist():
                data = zin.read(name)
                if name == "xl/worksheets/sheet2.xml":
                    data = build_data.swap_rows(data.decode(), a, b).encode()
                zout.writestr(name, data)
    return out.getvalue()


def corrupt_sheet2():
    """(flipped, bad_crc): sheet2.xml with one deflate byte inverted, or its CRC-32 changed in both headers."""
    data = bytearray(XLSX_BYTES)
    with zipfile.ZipFile(io.BytesIO(XLSX_BYTES)) as z:
        info = z.getinfo("xl/worksheets/sheet2.xml")
        cd_offset = z.start_dir
    off = info.header_offset
    name_len, extra_len = struct.unpack("<HH", data[off + 26:off + 30])
    start = off + 30 + name_len + extra_len
    flipped = bytearray(data)
    flipped[start + 50] ^= 0xFF
    bad = bytearray(data)
    bad[off + 14] ^= 0x01
    p = cd_offset
    while struct.unpack("<I", bad[p:p + 4])[0] == 0x02014B50:
        n, x, c = struct.unpack("<HHH", bad[p + 28:p + 34])
        if bad[p + 46:p + 46 + n] == b"xl/worksheets/sheet2.xml":
            bad[p + 16] ^= 0x01
        p += 46 + n + x + c
    return bytes(flipped), bytes(bad)


def _edit_sheet2(fn):
    with zipfile.ZipFile(io.BytesIO(XLSX_BYTES)) as zin:
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for name in zin.namelist():
                data = zin.read(name)
                if name == "xl/worksheets/sheet2.xml":
                    data = fn(data.decode()).encode()
                zout.writestr(name, data)
    return out.getvalue()


class TestRefusals(unittest.TestCase):
    def test_shuffled_rows_are_refused_with_both_rows_named(self):
        wb = LE.read_workbook(_swap_rows_in_snapshot(5, 10), name="shuffled.xlsx")
        with self.assertRaises(LE.LadderError) as cm:
            LE.analyze(wb, ROWS)
        msg = str(cm.exception)
        self.assertIn("ladder row 5 (construction_launch_per_kg): column A reads 'Boost price to customer ($/kg)', "
                      "expected a label containing 'launch price during construction'", msg)
        self.assertIn("ladder row 10 (customer_price_per_kg)", msg)
        self.assertIn("What would help", msg)

    def test_committed_fixture_is_that_swap(self):
        wb = LE.read_workbook(FIXTURE)
        self.assertEqual(wb["media"], [], "the fixture dropped the chart")
        with self.assertRaises(LE.LadderError) as cm:
            LE.analyze(wb, ROWS)
        self.assertIn("ladder row 5 (construction_launch_per_kg)", str(cm.exception))
        # row 8 is untouched: yellow is still B8:D8
        self.assertEqual(LE.census(wb["sheets"][1])["yellow_fill_cells"], ["B8", "C8", "D8"])

    def test_formula_in_a_ladder_cell_is_refused(self):
        bad = _edit_sheet2(lambda x: x.replace('<c r="B5" s="17" t="n"><v>2500</v></c>',
                                               '<c r="B5" s="17"><f>D5*3</f><v/></c>'))
        wb = LE.read_workbook(bad, name="formula.xlsx")
        with self.assertRaises(LE.LadderError) as cm:
            LE.analyze(wb, ROWS)
        self.assertIn("B5 (construction_launch_per_kg, Conservative) holds the formula =D5*3", str(cm.exception))

    def test_empty_ladder_cell_is_refused(self):
        bad = _edit_sheet2(lambda x: x.replace('<c r="C7" s="17" t="n"><v>0.06</v></c>', '<c r="C7" s="17" t="n"/>'))
        wb = LE.read_workbook(bad, name="empty.xlsx")
        with self.assertRaises(LE.LadderError) as cm:
            LE.analyze(wb, ROWS)
        self.assertIn("C7 (electricity_per_kwh, Baseline) is empty", str(cm.exception))

    def test_renamed_scenario_header_is_refused(self):
        bad = _edit_sheet2(lambda x: x.replace("<t>Baseline</t>", "<t>Nominal</t>"))
        wb = LE.read_workbook(bad, name="header.xlsx")
        with self.assertRaises(LE.LadderError) as cm:
            LE.analyze(wb, ROWS)
        self.assertIn("header C2 reads 'Nominal', expected the scenario column 'Baseline'", str(cm.exception))

    def test_missing_sheet_is_refused(self):
        rows = dict(ROWS, sheet="Inputs")
        with self.assertRaises(LE.LadderError) as cm:
            LE.analyze(WB, rows)
        self.assertIn("no sheet named 'Inputs'", str(cm.exception))

    def test_a_corrupt_member_is_a_workbook_error_not_a_crash(self):
        """sheet2.xml damaged under an intact central directory (a flipped deflate byte, a changed CRC)
        is 'cannot read', naming the member, never an uncaught zlib.error/BadZipFile."""
        flipped, bad_crc = corrupt_sheet2()
        with self.assertRaises(LE.WorkbookError) as cm:
            LE.read_workbook(flipped, name="flipped.xlsx")
        self.assertIn("member xl/worksheets/sheet2.xml could not be read", str(cm.exception))
        self.assertIn("decompressing", str(cm.exception))
        with self.assertRaises(LE.WorkbookError) as cm:
            LE.read_workbook(bad_crc, name="badcrc.xlsx")
        self.assertIn("member xl/worksheets/sheet2.xml could not be read: Bad CRC-32", str(cm.exception))
        with self.assertRaises(LE.WorkbookError):
            LE.read_workbook(XLSX_BYTES[:40000], name="truncated.xlsx")

    def test_not_a_zip_is_a_workbook_error(self):
        with self.assertRaises(LE.WorkbookError):
            LE.read_workbook(b"PK not really", name="x.xlsx")
        with self.assertRaises(LE.WorkbookError):
            LE.read_workbook(os.path.join(SNAPS, "rev_a.json"))


CASES = os.path.join(HERE, "fixtures", "malformed_cases.json")


def build_case(case: dict) -> bytes:
    """The case's base workbook with one member edited, re-zipped (the JS suite does the same)."""
    with open(os.path.join(HERE, _load(CASES)["bases"][case.get("base", "snapshot")]), "rb") as fh:
        base = fh.read()
    with zipfile.ZipFile(io.BytesIO(base)) as zin:
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for name in zin.namelist():
                data = zin.read(name)
                if name == case.get("member", "xl/worksheets/sheet2.xml"):
                    if "replace" in case:
                        old, new = (s.encode() for s in case["replace"])
                        assert old in data, case["name"]
                        data = data.replace(old, new, 1)
                    elif "truncate" in case:
                        data = data[:-case["truncate"]]
                    else:
                        after, byte = case["insert_byte"]
                        i = data.index(after.encode()) + len(after.encode())
                        data = data[:i] + bytes([byte]) + data[i:]
                zout.writestr(name, data)
    return out.getvalue()


def verdict(data: bytes):
    """('ok', result) | ('unreadable', message) | ('refused', message)."""
    try:
        return "ok", LE.analyze(LE.read_workbook(data, name="case.xlsx"), ROWS)
    except LE.WorkbookError as e:
        return "unreadable", str(e)
    except LE.LadderError as e:
        return "refused", str(e)


class TestCraftedDamage(unittest.TestCase):
    """fixtures/malformed_cases.json, shared with the JavaScript suite: both readers give each case the
    same verdict, and the CLI's exit code follows it."""

    def test_each_case_gets_its_verdict_and_message(self):
        cases = _load(CASES)["cases"]
        self.assertGreaterEqual(len(cases), 20)
        for case in cases:
            with self.subTest(case["name"]):
                kind, got = verdict(build_case(case))
                self.assertEqual(kind, case["expect"], got if isinstance(got, str) else "read without complaint")
                if kind == "ok":
                    key, scn, want = case["value"]
                    self.assertEqual(got["ladder"]["rows"][key]["values"][scn], want)
                    failed = [c["name"] for c in got["checks"] if not c["ok"]]
                    self.assertEqual(failed, [case["failedCheck"]] if "failedCheck" in case else [])
                else:
                    self.assertIn(case["message"], got)

    def test_the_cli_exit_code_follows_the_verdict_and_json_stays_json(self):
        import tempfile
        codes = {"ok": 0, "refused": 1, "unreadable": 2}
        with tempfile.TemporaryDirectory() as tmp:
            for i, case in enumerate(_load(CASES)["cases"]):
                with self.subTest(case["name"]):
                    path = os.path.join(tmp, f"case{i}.xlsx")
                    with open(path, "wb") as fh:
                        fh.write(build_case(case))
                    out, err = io.StringIO(), io.StringIO()
                    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
                        code = LE.main([path, "--format", "json"])
                    self.assertEqual(code, codes[case["expect"]], err.getvalue())
                    self.assertNotIn("Traceback", err.getvalue())
                    if code == 0:  # strict JSON: a bare NaN/Infinity raises
                        json.loads(out.getvalue(), parse_constant=lambda n: 1 / 0)
                    else:
                        self.assertEqual(out.getvalue(), "")

    def test_json_output_refuses_a_non_finite_number(self):
        with self.assertRaises(ValueError):
            LE.to_json({"x": float("nan")})
        with self.assertRaises(ValueError):
            LE.to_js({"x": float("inf")})


@functools.lru_cache(maxsize=None)
def recalculated():
    """The snapshot re-saved by headless LibreOffice (build_data.py): cached results, shared strings."""
    wb = LE.read_workbook(RECALC)
    res = LE.analyze(wb, ROWS)
    grid = {(s["name"], c["ref"]): c for s in wb["sheets"] for c in s["cells"]}
    return wb, res, grid


def cached_number(grid, sheet, ref):
    c = grid[(sheet, ref)]
    assert c["kind"] == "formula" and c["cached"] is not None, f"{sheet}!{ref} has no cached value"
    return float(c["cached"])


class TestRecalculatedCopy(unittest.TestCase):
    """A spreadsheet-saved copy: a shape this reader must handle, and an independent evaluation."""

    def test_census_shared_strings_and_the_same_ladder(self):
        wb, res, grid = recalculated()
        self.assertTrue(wb["application"].startswith("LibreOffice"), wb["application"])
        self.assertEqual(wb["shared_strings"], 158)
        self.assertEqual(wb["members"], 15, "the chart parts were removed before LibreOffice saved it")
        self.assertEqual(wb["media"], [])
        t = res["census"]["totals"]
        self.assertEqual((t["formula_cells"], t["formula_cells_with_cached_value"]), (146, 146))
        for name, c in res["census"]["sheets"].items():
            self.assertEqual(c["formula_cells_with_cached_value"], c["formula_cells"], name)
            self.assertEqual((c["literal_numbers"], c["strings"], c["blue_cells"]),
                             (RESULT["census"]["sheets"][name]["literal_numbers"], RESULT["census"]["sheets"][name]["strings"],
                              RESULT["census"]["sheets"][name]["blue_cells"]), name)
        self.assertIn("146 formula cells, all 146 with a cached result", res["finding"])
        self.assertNotEqual(res["workbook"]["md5"], RESULT["workbook"]["md5"])
        self.assertEqual([c["ok"] for c in res["checks"]], [True] * 4)
        for key, row in RESULT["ladder"]["rows"].items():
            self.assertEqual(res["ladder"]["rows"][key]["values"], row["values"], key)
            self.assertTrue(res["ladder"]["rows"][key]["blue"], key)
        self.assertEqual({k: v["value"] for k, v in res["design_point"]["entries"].items()},
                         {k: v["value"] for k, v in RESULT["design_point"]["entries"].items()})
        selector = grid[("Cash Flow & Payback", "B2")]
        self.assertEqual((selector["t"], selector["value"]), ("s", "Conservative"))

    def test_libreoffice_s_results_are_the_published_figures(self):
        _, _, grid = recalculated()
        mr = REV_A["mass_ratio"]["integrated_tether_to_payload_both_arms"]
        for col, k_ratio, k_tons, scn in (("B", "conservative_30pct_derate", "conservative_30pct_fiber_mass_t", "Conservative"),
                                          ("C", "baseline_15pct_derate", "baseline_15pct_fiber_mass_t", "Baseline"),
                                          ("D", "optimistic_0pct_derate", "optimistic_0pct_fiber_mass_t", "Optimistic")):
            self.assertAlmostEqual(2 * cached_number(grid, "Design Calcs", f"{col}11"), mr[k_ratio], places=2, msg=col)
            self.assertAlmostEqual(cached_number(grid, "Design Calcs", f"{col}12"), mr[k_tons], places=1, msg=col)
            mine = workbook_chain(scn)
            self.assertAlmostEqual(2 * cached_number(grid, "Design Calcs", f"{col}11"), mine["ratio"], places=9, msg=scn)
            self.assertAlmostEqual(cached_number(grid, "Design Calcs", f"{col}19") + ladder_value("program_fleet_capex_M", scn),
                                   mine["total_capex"], places=6, msg=scn)
        self.assertAlmostEqual(round(cached_number(grid, "Design Calcs", "B13"), 2), REV_A["kinematics"]["v_lower_tip_groundrel_kms"])
        m = re.search(r"Conservative [\d.]+× / \d+ t / \$([\d,]+)M / payback ~Y\d+", LOG)
        self.assertEqual(round(cached_number(grid, "Cash Flow & Payback", "B4")), int(m.group(1).replace(",", "")),
                         "the Cash Flow sheet's total capex with the Conservative selector")
        b16 = grid[("Cash Flow & Payback", "B16")]
        self.assertEqual((b16["t"], b16["cached"]), ("str", "beyond Y15"), "payback cell is a string result")
        self.assertLess(cached_number(grid, "Cash Flow & Payback", "P15"), 0, "cumulative cash at Y15 is negative")
        # B15:P15 (Conservative selector) against the transcription
        mine = workbook_chain("Conservative")["cumulative"]
        for i, col in enumerate("BCDEFGHIJKLMNOP"):
            self.assertAlmostEqual(cached_number(grid, "Cash Flow & Payback", f"{col}15"), mine[i], places=6, msg=col)
        pb = [grid[("Phase Budgets", f"{c}{r}")]["value"] for c in "BC" for r in range(3, 7)]
        self.assertEqual(cached_number(grid, "Phase Budgets", "B7"), sum(pb[:4]))
        self.assertEqual(cached_number(grid, "Phase Budgets", "C7"), sum(pb[4:]))

    def test_erf_audit_block_fails_under_spreadsheet_precedence(self):
        """Finding 6: rows 28-31 check ERF() against A&S 7.1.26, "should be < 2e-7". EXP(-B9^2) binds
        as EXP((-B9)^2) in spreadsheets, so the check fails by about 5.8. A fixed workbook fails this."""
        _, _, grid = recalculated()
        x = cached_number(grid, "Design Calcs", "B9")
        t = cached_number(grid, "Design Calcs", "B28")
        self.assertAlmostEqual(t, 1 / (1 + 0.3275911 * x), places=9)
        formula = grid[("Design Calcs", "B29")]["formula"]
        self.assertIn("EXP(-B9^2)", formula)
        terms = re.findall(r"([-+]?\d+\.\d+)\*B28(?:\^(\d))?", formula)
        self.assertEqual(len(terms), 5)
        poly = sum(float(a) * t ** (int(p) if p else 1) for a, p in terms)
        intended = 1 - poly * math.exp(-x * x)
        as_evaluated = 1 - poly * math.exp(+x * x)
        self.assertLess(abs(intended - math.erf(x)), 2e-7, "the constants are A&S 7.1.26 and the intent is right")
        self.assertAlmostEqual(cached_number(grid, "Design Calcs", "B29"), as_evaluated, places=6)
        self.assertAlmostEqual(cached_number(grid, "Design Calcs", "B30"), math.erf(x), places=9)
        self.assertGreater(cached_number(grid, "Design Calcs", "B31"), 1.0)
        self.assertIn("should be < 2e-7", grid[("Design Calcs", "E31")]["value"])
        for col in "CD":
            self.assertGreater(cached_number(grid, "Design Calcs", f"{col}31"), 1.0, col)


@functools.lru_cache(maxsize=None)
def pre_fix():
    wb = LE.read_workbook(PRE_FIX)
    grid = {(s["name"], c["ref"]): c for s in wb["sheets"] for c in s["cells"]}
    return wb, grid


class TestPreFixWorkbook(unittest.TestCase):
    """The workbook before the ladder fix (blob 46bc02e at 7bb5b5f, parent of the fix 2c27b05); the
    local Skyhook checkout is shallow (from 30419ef), so it is not there."""

    def test_blob_id_is_the_one_github_reports(self):
        with open(PRE_FIX, "rb") as fh:
            data = fh.read()
        self.assertEqual(hashlib.sha1(b"blob %d\0" % len(data) + data).hexdigest(),
                         "46bc02ed26367805f0e920b00a52de9098e42613")
        # the same function gives the snapshot its sidecar blob
        self.assertEqual(hashlib.sha1(b"blob %d\0" % len(XLSX_BYTES) + XLSX_BYTES).hexdigest(), PROV["committedBlob"])

    def test_it_stored_134_results_written_by_libreoffice(self):
        wb, _ = pre_fix()
        self.assertTrue(wb["application"].startswith("LibreOffice/24.2.7.2"), wb["application"])
        per = {s["name"]: LE.census(s) for s in wb["sheets"]}
        self.assertEqual({n: (c["formula_cells"], c["formula_cells_with_cached_value"]) for n, c in per.items()},
                         {"README": (0, 0), "Assumptions": (0, 0), "Design Calcs": (69, 69),
                          "Phase Budgets": (2, 2), "Cash Flow & Payback": (63, 63)})
        self.assertEqual(LE.totals(per)["formula_cells"], 134, "docs/original_README.md:29's '134 live formulas'")

    def test_the_twelve_formulas_added_since_are_the_erf_audit_block(self):
        wb, _ = pre_fix()

        def formulas(w, name):
            s = next(s for s in w["sheets"] if s["name"] == name)
            return {c["ref"]: c["formula"] for c in s["cells"] if c["kind"] == "formula"}
        for name in ("Design Calcs", "Phase Budgets", "Cash Flow & Payback"):
            old, new = formulas(wb, name), formulas(WB, name)
            added = sorted(set(new) - set(old))
            self.assertEqual(set(old) - set(new), set(), name)
            self.assertEqual({r: new[r] for r in old}, old, f"{name}: every pre-fix formula is unchanged")
            self.assertEqual(added, [f"{c}{r}" for c in "BCD" for r in range(28, 32)] if name == "Design Calcs" else [], name)
        self.assertFalse(any(c["kind"] == "string" and "ERF" in (c["value"] or "") for s in wb["sheets"] for c in s["cells"]))

    def test_the_a18_note_was_true_before_the_fix_and_was_kept_after_it(self):
        _, grid = pre_fix()
        now = {(s["name"], c["ref"]): c for s in WB["sheets"] for c in s["cells"]}
        note = grid[("Cash Flow & Payback", "A18")]["value"]
        self.assertIn("Conservative scenario also spins slower (row 15), per Rule 2.3.", note)
        self.assertEqual(now[("Cash Flow & Payback", "A18")]["value"], note, "unchanged by the fix")
        self.assertIn("conservative spins slower", grid[("Assumptions", "A15")]["value"])
        self.assertEqual([grid[("Assumptions", f"{c}15")]["value"] for c in "BCD"], [2.0, 2.8, 3.5])
        self.assertEqual([ladder_value("tip_speed_kms", s) for s in SCN], [2.5, 2.5, 2.5])

    def test_its_stored_results_are_the_formula_chain_on_its_own_inputs(self):
        """A second, older LibreOffice evaluation, on different inputs, against the transcription."""
        wb, grid = pre_fix()
        cols = ROWS["columns"]
        dp = LE.extract_block(wb, "Assumptions", LE.DESIGN_POINT_HEADER, (LE.CONSTANTS_HEADER, LE.LADDER_HEADER))["entries"]
        cs = LE.extract_block(wb, "Assumptions", LE.CONSTANTS_HEADER, (LE.DESIGN_POINT_HEADER, LE.LADDER_HEADER))["entries"]
        got = workbook_chain("Conservative",
                             inputs=lambda k, s: grid[("Assumptions", f"{cols[s]}{ROWS['rows'][k]}")]["value"],
                             design=lambda k: dp[k]["value"], constant=lambda k: cs[k]["value"])
        self.assertEqual(grid[("Cash Flow & Payback", "B2")]["value"], "Conservative")
        self.assertAlmostEqual(got["ratio"], 2 * cached_number(grid, "Design Calcs", "B11"), places=9)
        self.assertAlmostEqual(got["total_capex"], cached_number(grid, "Cash Flow & Payback", "B4"), places=6)
        self.assertEqual(got["payback"], int(cached_number(grid, "Cash Flow & Payback", "B16")))
        self.assertEqual(got["payback"], 15)
        self.assertEqual(round(got["total_capex"]), 7085, "not the post-fix $8,741M: the file never stored those")

    def test_the_row_guard_refuses_it_on_row_3(self):
        wb, _ = pre_fix()
        with self.assertRaises(LE.LadderError) as cm:
            LE.analyze(wb, ROWS)
        msg = str(cm.exception)
        self.assertIn("ladder row 3 (strength_retention): column A reads 'Fiber strength derate (joints/weave/aging)', "
                      "expected a label containing 'strength retention'", msg)
        self.assertEqual(msg.count("ladder row"), 1, "only row 3's label changed in the fix")


class TestOutputs(unittest.TestCase):
    def test_json_is_stable_and_the_committed_artefact_is_current(self):
        a = LE.to_json(LE.analyze(LE.read_workbook(XLSX), ROWS, PROV))
        b = LE.to_json(LE.analyze(LE.read_workbook(XLSX), ROWS, PROV))
        self.assertEqual(a, b)
        with open(LADDER_JSON, encoding="utf-8") as fh:
            committed = fh.read()
        self.assertEqual(committed, a, "snapshots/ladder.json is what the tool writes today")
        self.assertEqual(a.count("copiedAt"), 1, "no timestamp of its own; the only date is the embedded sidecar's")
        without_prov = LE.to_json(LE.analyze(LE.read_workbook(XLSX), ROWS))
        self.assertNotIn("copiedAt", without_prov)
        self.assertNotIn(str(2026), without_prov.split('"created"')[0], "nothing dated before the workbook's own docProps")

    def test_page_data_is_the_same_json(self):
        with open(os.path.join(HERE, "page-data.js"), encoding="utf-8") as fh:
            js = fh.read()
        body = js[js.index("window.BLUE_CELLS_DATA = ") + len("window.BLUE_CELLS_DATA = "):].rstrip().rstrip(";")
        self.assertEqual(json.loads(body), json.loads(LE.to_json(RESULT)))

    def test_report_lines(self):
        rep = LE.render_report(RESULT)
        self.assertIn("Ladder    Assumptions!B3:D16  14 rows x 3 columns", rep)
        self.assertIn("  construction_launch_per_kg            2500          1500           800  B5:D5    "
                      "Launch price during construction ($/kg)", rep)
        self.assertIn("  TOTAL                     395       146       0       81      164    63      5       4", rep)
        self.assertIn("Checks    4 of 4 passed (reported, not enforced)", rep)
        self.assertIn("  held ladder-rows: 14 rows from finance/skyhook_finance.py found where the map says", rep)
        self.assertIn("provenance  https://github.com/hudsonclavin-cloud/skyhook tools/financial/skyhook_financial_model.xlsx @ 34935cb", rep)


class TestCommandLine(unittest.TestCase):
    def run_cli(self, *args):
        return subprocess.run([sys.executable, "-I", os.path.join(HERE, "ladder_export.py"), *args],
                              capture_output=True, text=True, cwd=HERE)

    def test_help_and_version(self):
        r = self.run_cli("--help")
        self.assertEqual(r.returncode, 0)
        self.assertIn("Exit 1: refused", r.stdout)
        self.assertIn("--rows-map", r.stdout)
        self.assertEqual(self.run_cli("--version").stdout.strip(), f"ladder_export.py {LE.__version__}")

    def test_default_run_is_the_snapshot(self):
        r = self.run_cli()
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("md5     5ccdee4d43e91f8dd98884678c8e99bb", r.stdout)
        self.assertIn("146       0", r.stdout)

    def test_json_format_and_out_and_js(self):
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "l.json")
            js = os.path.join(tmp, "l.js")
            r = self.run_cli("--out", out, "--js", js, "--format", "json")
            self.assertEqual(r.returncode, 0, r.stderr)
            printed = json.loads(r.stdout)
            with open(out, encoding="utf-8") as fh:
                written = json.load(fh)
            self.assertEqual(printed, written)
            self.assertEqual(written["census"]["totals"]["formula_cells"], 146)
            with open(js, encoding="utf-8") as fh:
                text = fh.read()
            self.assertTrue(text.startswith("// Written by ladder_export.py --js"))
            self.assertIn("window.BLUE_CELLS_DATA = {", text)

    def test_exit_codes(self):
        r = self.run_cli(FIXTURE)
        self.assertEqual(r.returncode, 1)
        self.assertIn("REFUSED", r.stderr)
        self.assertIn("ladder row 5 (construction_launch_per_kg): column A reads 'Boost price to customer ($/kg)'", r.stderr)
        self.assertEqual(r.stdout, "")
        self.assertEqual(self.run_cli("/nonexistent/workbook.xlsx").returncode, 2)
        r = self.run_cli(os.path.join(SNAPS, "rev_a.json"))
        self.assertEqual(r.returncode, 2)
        self.assertIn("not a zip archive", r.stderr)
        self.assertEqual(self.run_cli("--rows-map", "/nonexistent.json").returncode, 2)
        r = self.run_cli("--rows-map", os.path.join(SNAPS, "decision_grid.json"))
        self.assertEqual((r.returncode, "Traceback" in r.stderr), (2, False), r.stderr)
        self.assertIn("the row map must hold", r.stderr)

    def test_a_corrupt_member_exits_2_without_a_traceback(self):
        import tempfile
        flipped, bad_crc = corrupt_sheet2()
        with tempfile.TemporaryDirectory() as tmp:
            for label, data in (("flipped", flipped), ("badcrc", bad_crc)):
                path = os.path.join(tmp, f"{label}.xlsx")
                with open(path, "wb") as fh:
                    fh.write(data)
                r = self.run_cli(path)
                self.assertEqual(r.returncode, 2, f"{label}: {r.stderr}")
                self.assertNotIn("Traceback", r.stderr, label)
                self.assertIn("cannot read the workbook — member xl/worksheets/sheet2.xml could not be read", r.stderr)
                self.assertEqual(r.stdout, "")

    def test_an_internal_error_exits_3_not_1(self):
        def broken(*_a, **_k):
            raise KeyError("a bug")
        err = io.StringIO()
        original = LE.analyze
        LE.analyze = broken
        try:
            with contextlib.redirect_stderr(err), contextlib.redirect_stdout(io.StringIO()):
                code = LE.main([XLSX])
        finally:
            LE.analyze = original
        self.assertEqual(code, 3)
        self.assertIn("internal error", err.getvalue())
        self.assertIn("KeyError", err.getvalue())

    def test_an_output_it_cannot_write_exits_2_without_a_traceback(self):
        """A missing directory, a directory, a closed stdout, a bad sidecar or row map: exit 2, one
        line, never exit 1 (a refusal)."""
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            missing = os.path.join(tmp, "no_such_dir", "ladder.json")
            r = self.run_cli("--out", missing)
            self.assertEqual((r.returncode, "Traceback" in r.stderr), (2, False), r.stderr)
            self.assertEqual(r.stderr, f"ladder_export.py: cannot write {missing} — [Errno 2] No such file or directory: "
                                       f"'{missing}'\n")
            self.assertEqual(r.stdout, "")
            r = self.run_cli("--js", tmp)
            self.assertEqual((r.returncode, "Traceback" in r.stderr), (2, False), r.stderr)
            self.assertTrue(r.stderr.startswith(f"ladder_export.py: cannot write {tmp} — "), r.stderr)
            self.assertEqual(r.stderr.count("\n"), 1)
            prov = os.path.join(tmp, "list.json")
            with open(prov, "w", encoding="utf-8") as fh:
                fh.write("[1, 2]")
            r = self.run_cli("--provenance", prov)
            self.assertEqual((r.returncode, "Traceback" in r.stderr), (2, False), r.stderr)
            self.assertIn("the provenance sidecar must be a JSON object", r.stderr)
            rows = os.path.join(tmp, "rows.json")
            with open(rows, "w", encoding="utf-8") as fh:
                json.dump(dict(ROWS, scenarios=[["Conservative"]]), fh)
            r = self.run_cli("--rows-map", rows)
            self.assertEqual((r.returncode, "Traceback" in r.stderr), (2, False), r.stderr)
        # stdout whose reader has gone (EPIPE)
        rd, wr = os.pipe()
        os.close(rd)
        try:
            p = subprocess.run([sys.executable, "-I", os.path.join(HERE, "ladder_export.py")],
                               stdout=wr, stderr=subprocess.PIPE, text=True, cwd=HERE)
        finally:
            os.close(wr)
        self.assertEqual((p.returncode, "Traceback" in p.stderr), (2, False), p.stderr)
        self.assertEqual(p.stderr, "ladder_export.py: cannot write standard output — [Errno 32] Broken pipe\n")

    def test_an_internal_error_while_writing_exits_3_not_1(self):
        def broken(*_a, **_k):
            raise KeyError("a bug in to_js")
        import tempfile
        err = io.StringIO()
        original = LE.to_js
        LE.to_js = broken
        try:
            with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stderr(err), \
                    contextlib.redirect_stdout(io.StringIO()):
                code = LE.main([XLSX, "--js", os.path.join(tmp, "x.js")])
        finally:
            LE.to_js = original
        self.assertEqual(code, 3)
        self.assertIn("a bug in to_js", err.getvalue())

    def test_census_reads_a_file_the_row_map_refuses(self):
        r = self.run_cli(PRE_FIX)
        self.assertEqual(r.returncode, 1)
        r = self.run_cli("--census", PRE_FIX)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("  TOTAL                     372       134     134       81      153    63      5       4", r.stdout)
        self.assertIn("This workbook holds 134 formula cells, all 134 with a cached result", r.stdout)
        self.assertNotIn("Ladder", r.stdout)

    def test_example_txt_matches_a_live_run_when_the_checkout_is_present(self):
        with open(os.path.join(HERE, "example.txt"), encoding="utf-8") as fh:
            example = fh.read()
        blocks = re.split(r"^(?=\$ )", example, flags=re.M)
        first = next(b for b in blocks if b.startswith("$ "))
        cmd_line, _, body = first.partition("\n")
        body = body.rsplit("(exit", 1)[0]
        self.assertIn("ladder_export.py", cmd_line)
        head = subprocess.run(["git", "rev-parse", "HEAD"], cwd=SKYHOOK, capture_output=True, text=True) \
            if os.path.isdir(os.path.join(SKYHOOK, ".git")) else None
        if head is None or head.returncode != 0 or head.stdout.strip() != PROV["commit"]:
            self.skipTest(f"{SKYHOOK} is not at {PROV['commit'][:7]}; the saved example is compared only against "
                          "the checkout it was made from")
        args = cmd_line.split()[1:]
        args = [a.replace("ideas/08-blue-cells/", HERE + "/") for a in args]
        r = subprocess.run([sys.executable, "-I"] + args[1:] if args and args[0] == "python3" else args,
                           capture_output=True, text=True, cwd=os.path.join(HERE, "..", ".."))
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(r.stdout, body)


if __name__ == "__main__":
    unittest.main()
